# -*- coding: utf-8 -*-
"""把 命运牌局_配置总表.xlsx 导出为 data/config.json，供网页读取。"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "命运牌局_配置总表.xlsx"
OUT_DIR = ROOT.parent / "data"
OUT_JSON = OUT_DIR / "config.json"


def sheet_to_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = []
    for h in rows[0]:
        if h is None:
            headers.append("")
        else:
            headers.append(str(h).strip())
    # stop at first fully empty header run
    while headers and headers[-1] == "":
        headers.pop()
    out = []
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row[: len(headers)]):
            continue
        # skip section title rows (first cell looks like a second header block)
        first = row[0]
        if first is not None and str(first).strip() in headers:
            continue
        item = {}
        empty = True
        for i, key in enumerate(headers):
            if not key:
                continue
            val = row[i] if i < len(row) else None
            if val is None:
                item[key] = ""
            elif isinstance(val, float) and val.is_integer():
                item[key] = int(val)
                empty = False
            else:
                item[key] = val
                if str(val).strip() != "":
                    empty = False
        if not empty:
            out.append(item)
    return out


def constants_to_map(rows):
    m = {}
    for r in rows:
        key = str(r.get("常量键", "")).strip()
        if not key:
            continue
        val = r.get("数值/文本", "")
        m[key] = val
    return m


def main():
    if not XLSX.exists():
        raise SystemExit(f"missing xlsx: {XLSX}")
    wb = load_workbook(XLSX, data_only=True)
    config = {
        "_meta": {
            "source": XLSX.name,
            "sheets": wb.sheetnames,
        },
        "sheets": {},
    }
    for name in wb.sheetnames:
        if name.startswith("00_"):
            continue
        ws = wb[name]
        rows = sheet_to_dicts(ws)
        # normalize key: 01_效果类型说明 -> effects_doc style keep Chinese name
        config["sheets"][name] = rows

    # convenience maps
    const_rows = config["sheets"].get("02_通用常量", [])
    config["constants"] = constants_to_map(const_rows)
    config["relics"] = {r.get("宝物ID"): r for r in config["sheets"].get("06_宝物", []) if r.get("宝物ID")}
    config["heritages"] = {r.get("传承ID"): r for r in config["sheets"].get("05_传承卡", []) if r.get("传承ID")}
    config["fateCards"] = {r.get("命运卡ID"): r for r in config["sheets"].get("07_命运卡", []) if r.get("命运卡ID")}
    config["monsters"] = {r.get("怪物ID"): r for r in config["sheets"].get("08_怪物", []) if r.get("怪物ID")}
    config["companions"] = {r.get("同伴ID"): r for r in config["sheets"].get("17_同伴", []) if r.get("同伴ID")}
    config["warehouseItems"] = {
        r.get("物品ID"): r for r in config["sheets"].get("26_仓库物品", []) if r.get("物品ID")
    }

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", OUT_JSON)
    print(
        "counts",
        {
            "relics": len(config["relics"]),
            "heritages": len(config["heritages"]),
            "fateCards": len(config["fateCards"]),
            "monsters": len(config["monsters"]),
            "constants": len(config["constants"]),
        },
    )


if __name__ == "__main__":
    main()
